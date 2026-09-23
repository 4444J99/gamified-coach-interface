"""
Unit tests for DocumentAnalyzer in scripts/analyze_docs.py
"""

import os
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from docx import Document

# Ensure scripts directory is in sys.path
sys.path.insert(0, str(Path(__file__).parent.parent))

from analyze_docs import DocumentAnalyzer  # noqa: E402


class TestDocumentAnalyzer(unittest.TestCase):
    """Test suite for DocumentAnalyzer class."""

    def setUp(self):
        """Set up temporary directory and sample documents for testing."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)

        self.valid_docx_path = self.base_path / "valid_test.docx"
        doc = Document()
        doc.add_heading("Fitness Blueprint", level=1)
        doc.add_paragraph("This is a fitness training and workout exercise plan.")
        doc.add_paragraph("It provides coaching strategies and business models.")
        doc.save(self.valid_docx_path)

    def tearDown(self):
        """Clean up temporary directory."""
        self.temp_dir.cleanup()

    def _capture_ingest(self, analyzer):
        captured_output = StringIO()
        with redirect_stdout(captured_output):
            ingested = analyzer.ingest_docs()
        return ingested, captured_output.getvalue()

    def test_ingest_and_digest_valid_docs(self):
        """Test ingesting and digesting valid docx files."""
        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        ingested = analyzer.ingest_docs()

        self.assertIn("valid_test.docx", ingested)
        self.assertIn("valid_test.docx", analyzer.documents)
        self.assertEqual(analyzer.documents["valid_test.docx"]["num_paragraphs"], 3)

        analysis = analyzer.digest_docs()
        self.assertEqual(analysis["total_documents"], 1)
        self.assertIn("valid_test.docx", analysis["documents"])
        doc_analysis = analysis["documents"]["valid_test.docx"]
        self.assertGreater(doc_analysis["word_count"], 0)
        self.assertTrue(len(doc_analysis["key_topics"]) > 0)

    def test_oversized_file_rejection(self):
        """Test that files exceeding max_file_size_bytes are rejected."""
        analyzer = DocumentAnalyzer(base_dir=str(self.base_path), max_file_size_bytes=10)
        ingested, output = self._capture_ingest(analyzer)

        self.assertNotIn("valid_test.docx", ingested)
        self.assertIn("File size exceeds limit", output)

    def test_expanded_archive_size_is_bounded_before_parser(self):
        """Highly compressed DOCX content is rejected by expanded-size limit."""
        bomb_path = self.base_path / "compressed-bomb.docx"
        bomb_doc = Document()
        bomb_doc.add_paragraph("small visible document")
        bomb_doc.save(bomb_path)
        with zipfile.ZipFile(bomb_path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("customXml/highly-compressible.xml", "A" * 200_000)

        self.assertLess(bomb_path.stat().st_size, 100_000)
        analyzer = DocumentAnalyzer(
            base_dir=str(self.base_path),
            max_file_size_bytes=100_000,
            max_expanded_file_size_bytes=100_000,
        )
        ingested, output = self._capture_ingest(analyzer)

        self.assertNotIn("compressed-bomb.docx", ingested)
        self.assertIn("Expanded document size exceeds limit", output)

    def test_symlink_outside_base_dir_is_rejected(self):
        """A .docx symlink must not escape the configured source root."""
        with tempfile.TemporaryDirectory() as outside_temp_dir:
            outside_path = Path(outside_temp_dir) / "outside.docx"
            outside_doc = Document()
            outside_doc.add_paragraph("outside source root")
            outside_doc.save(outside_path)

            linked_path = self.base_path / "linked-outside.docx"
            try:
                os.symlink(outside_path, linked_path)
            except (OSError, NotImplementedError) as exc:
                self.skipTest(f"symlinks unavailable in this environment: {exc}")

            analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
            ingested, output = self._capture_ingest(analyzer)

            self.assertNotIn("linked-outside.docx", ingested)
            self.assertIn("linked-outside.docx", output)
            self.assertIn("Access denied (path outside base directory)", output)

    def test_symlink_loop_is_rejected_without_aborting_scan(self):
        """An unresolvable .docx symlink is rejected while valid files continue."""
        loop_path = self.base_path / "loop.docx"
        try:
            os.symlink(loop_path, loop_path)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"symlinks unavailable in this environment: {exc}")

        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        ingested, output = self._capture_ingest(analyzer)

        self.assertIn("valid_test.docx", ingested)
        self.assertNotIn("loop.docx", ingested)
        self.assertIn("Unable to resolve file path", output)

    def test_hidden_parent_of_base_does_not_filter_documents(self):
        """Filtering applies only to descendants of base_dir, not its ancestors."""
        hidden_workspace = self.base_path / ".workspace"
        project_root = hidden_workspace / "project"
        project_root.mkdir(parents=True)
        visible_docx = project_root / "visible.docx"
        doc = Document()
        doc.add_paragraph("visible")
        doc.save(visible_docx)

        analyzer = DocumentAnalyzer(base_dir=str(project_root))
        ingested, _ = self._capture_ingest(analyzer)
        self.assertIn("visible.docx", ingested)

    def test_non_regular_docx_path_is_rejected(self):
        """A directory with a .docx suffix must not be parsed as a document."""
        fake_docx_dir = self.base_path / "directory.docx"
        fake_docx_dir.mkdir()

        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        ingested, output = self._capture_ingest(analyzer)

        self.assertNotIn("directory.docx", ingested)
        self.assertIn("Path is not a regular file", output)

    def test_parser_receives_stable_open_file_handle(self):
        """Parser consumes the already-open validated handle rather than reopening a path."""
        from docx import Document as RealDocument

        parser_sources = []

        def parse_from_handle(source):
            parser_sources.append(source)
            return RealDocument(source)

        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        with patch("analyze_docs.Document", side_effect=parse_from_handle):
            ingested, _ = self._capture_ingest(analyzer)

        self.assertIn("valid_test.docx", ingested)
        self.assertTrue(parser_sources)
        self.assertTrue(hasattr(parser_sources[0], "read"))
        self.assertNotIsInstance(parser_sources[0], (str, Path))

    def test_corrupted_file_handling(self):
        """Corrupted docx files fail safely without a traceback."""
        corrupt_path = self.base_path / "corrupt.docx"
        with open(corrupt_path, "wb") as corrupt_file:
            corrupt_file.write(b"NOT_A_REAL_DOCX_FILE_DATA_1234567890")

        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        ingested, output = self._capture_ingest(analyzer)

        self.assertNotIn("corrupt.docx", ingested)
        self.assertIn("Failed to ingest corrupt.docx", output)
        self.assertNotIn("Traceback", output)

    def test_unexpected_parser_error_does_not_leak_exception_message(self):
        """Unexpected parser failures expose the class, not raw sensitive details."""
        sensitive_detail = "/private/customer/path/token-like-detail"
        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))

        with patch(
            "analyze_docs.Document",
            side_effect=RuntimeError(sensitive_detail),
        ):
            ingested, output = self._capture_ingest(analyzer)

        self.assertNotIn("valid_test.docx", ingested)
        self.assertIn("Ingestion error (RuntimeError)", output)
        self.assertNotIn(sensitive_detail, output)

    def test_stdout_capture_is_restored(self):
        """Ingestion tests preserve an enclosing runner's stdout stream."""
        outer_stream = StringIO()
        original_stdout = sys.stdout
        try:
            sys.stdout = outer_stream
            analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
            self._capture_ingest(analyzer)
            self.assertIs(sys.stdout, outer_stream)
        finally:
            sys.stdout = original_stdout

    def test_key_topics_extraction(self):
        """Test topic extraction algorithm."""
        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        paragraphs = [
            "Fitness fitness fitness training workout exercise.",
            "Gamified gamified battle quest gaming fitness.",
        ]
        topics = analyzer._extract_key_topics(paragraphs)
        self.assertIsInstance(topics, list)
        self.assertIn("fitnes", topics)  # Stemmed fitness

    def test_summary_creation_truncation(self):
        """Test summary creation and truncation behavior."""
        analyzer = DocumentAnalyzer(base_dir=str(self.base_path))
        long_paragraph = "word " * 100
        summary = analyzer._create_summary([long_paragraph])
        self.assertTrue(summary.endswith("..."))
        self.assertLessEqual(len(summary), 300)


if __name__ == "__main__":
    unittest.main()
